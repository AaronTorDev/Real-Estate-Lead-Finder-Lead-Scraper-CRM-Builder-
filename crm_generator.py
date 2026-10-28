from playwright.sync_api import sync_playwright
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

directory_url = "https://www.camara.cr/directorio-de-asociados/"
output_file = "crm_prospects_complete.xlsx"

print("Starting total directory extraction...")

filtered_prospects = []

with sync_playwright() as p:
    print("Launching browser...")
    browser = p.chromium.launch(headless=True)
    context = browser.new_context(viewport={"width": 1280, "height": 1080})
    page = context.new_page()
    
    print(f"Opening directory: {directory_url}")
    page.goto(directory_url, timeout=90000)
    page.wait_for_timeout(5000)
    
    scroll_position = 0
    empty_count = 0
    max_attempts_without_change = 8
    
    print("Scanning the entire directory...")
    
    while True:
        scroll_position += 1500
        page.evaluate(f"window.scrollTo(0, {scroll_position});")
        page.wait_for_timeout(2000)
        
        current_html = page.content()
        soup = BeautifulSoup(current_html, 'html.parser')
        
        broker_blocks = soup.find_all('div', class_='jet-listing-grid__item')
        if not broker_blocks:
            broker_blocks = [t for t in soup.find_all('div', class_='elementor-widget-wrap elementor-element-populated') 
                             if t.find('div', class_='jet-listing-dynamic-field__content')]

        before_adding = len(filtered_prospects)
        
        for block in broker_blocks:
            text_fields = block.find_all('div', class_='jet-listing-dynamic-field__content')
            texts = [c.text.strip() for c in text_fields if c.text]
            
            if not texts or len(texts[0]) < 3:
                continue
                
            if len(texts) > 1 and "Corredor" not in texts[1] and "Tel." not in texts[1] and len(texts[1]) < 30:
                full_name = f"{texts[0]} {texts[1]}".upper()
            else:
                full_name = texts[0].upper()
            
            if "BUSCAR" in full_name or "BIENVENIDO" in full_name or "REALTOR" in full_name:
                continue
                
            if full_name in [p['Name'] for p in filtered_prospects]:
                continue

            link_widgets = block.find_all('div', class_='jet-listing-dynamic-link')
            
            fb_link, ig_link, li_link, web_link = "", "", "", ""
            has_fb, has_ig, has_web = False, False, False
            base_points = 0

            for widget in link_widgets:
                tag_a = widget.find('a', href=True)
                tag_i = widget.find('i', class_=True)
                
                if not tag_a or not tag_i:
                    continue
                    
                href = tag_a['href']
                icon_classes = " ".join(tag_i['class']).lower()
                
                if "fa-facebook" in icon_classes or "facebook.com" in href.lower():
                    fb_link = href
                    has_fb = True
                    base_points += 2
                elif "fa-instagram" in icon_classes or "instagram.com" in href.lower():
                    ig_link = href
                    has_ig = True
                    base_points += 2
                elif "fa-linkedin" in icon_classes or "linkedin.com" in href.lower():
                    li_link = href
                    base_points += 1
                elif "fa-globe" in icon_classes or ("http" in href.lower() and not any(x in href.lower() for x in ["wa.me", "whatsapp", "camara.cr"])):
                    web_link = href
                    has_web = True
                    base_points += 4

            extra_points = 0
            if has_fb and has_ig and has_web:
                extra_points = 2.0
            elif has_ig and has_web:
                extra_points = 1.0
            elif has_fb and has_web:
                extra_points = 0.5
                
            final_score = base_points + extra_points

            if final_score == 0:
                continue

            phone = "Not specified"
            for t in texts:
                if 'Tel. celular:' in t:
                    phone = t.replace('Tel. celular:', '').strip()
                    
            email_tag = block.find('a', href=lambda h: h and h.startswith('mailto:'))
            email = "Not specified"
            if email_tag:
                dirty_email = email_tag.get('href').replace('mailto:', '').strip()
                email = dirty_email.split('%20')[0].strip() if "%20" in dirty_email else dirty_email
            
            filtered_prospects.append({
                'Name': full_name,
                'Phone': phone,
                'Email': email,
                'Web': web_link,
                'IG': ig_link,
                'FB': fb_link,
                'LinkedIn': li_link,
                'Score': final_score
            })
        
        after_adding = len(filtered_prospects)
        print(f"    [Progress] Accumulated total: {after_adding} valid prospects...")
        
        if after_adding == before_adding:
            empty_count += 1
            if empty_count >= max_attempts_without_change:
                print("Reached the actual end of the directory.")
                break
        else:
            empty_count = 0
            
    browser.close()

print("\nGenerating master CRM file...")

filtered_prospects.sort(key=lambda x: x['Score'], reverse=True)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "CRM Leads"
ws.views.sheetView[0].showGridLines = True

headers = [
    "ID", "SCORE", "FULL NAME", "WEB", "IG", "FACEBOOK", "LINKEDIN", 
    "EMAIL", "PHONE", "REVIEW STATUS", "WORK STATUS", "NOTES"
]
ws.append(headers)

font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center")

for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

review_options = '"[ ] Unreviewed,[X] Reviewed"'
work_options = '"[ ] To Evaluate,REJECTED,PENDING (Free sample),FINISHED (Contact ready)"'

dv_review = DataValidation(type="list", formula1=review_options, allow_blank=True)
dv_work = DataValidation(type="list", formula1=work_options, allow_blank=True)

ws.add_data_validation(dv_review)
ws.add_data_validation(dv_work)

for idx, p in enumerate(filtered_prospects, 1):
    row_num = idx + 1
    
    ws.cell(row=row_num, column=1, value=idx).alignment = align_center
    ws.cell(row=row_num, column=2, value=p['Score']).alignment = align_center
    ws.cell(row=row_num, column=3, value=p['Name']).alignment = align_left
    
    for col_idx, key in [(4, 'Web'), (5, 'IG'), (6, 'FB'), (7, 'LinkedIn')]:
        cell = ws.cell(row=row_num, column=col_idx)
        if p[key]:
            cell.value = "Open Link"
            cell.hyperlink = p[key]
            cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        else:
            cell.value = "-"
        cell.alignment = align_center
        
    ws.cell(row=row_num, column=8, value=p['Email']).alignment = align_left
    ws.cell(row=row_num, column=9, value=p['Phone']).alignment = align_center
    
    review_cell = ws.cell(row=row_num, column=10, value="[ ] Unreviewed")
    review_cell.alignment = align_center
    dv_review.add(review_cell)
    
    work_cell = ws.cell(row=row_num, column=11, value="[ ] To Evaluate")
    work_cell.alignment = align_center
    dv_work.add(work_cell)
    
    ws.cell(row=row_num, column=12, value="").alignment = align_left

review_range = f"J2:J{len(filtered_prospects) + 1}"
work_range = f"K2:K{len(filtered_prospects) + 1}"

fill_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
fill_red = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
fill_blue = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

ws.conditional_formatting.add(review_range, CellIsRule(operator='equal', formula=['"[X] Reviewed"'], fill=fill_green))
ws.conditional_formatting.add(work_range, CellIsRule(operator='equal', formula=['"REJECTED"'], fill=fill_red))
ws.conditional_formatting.add(work_range, CellIsRule(operator='equal', formula=['"PENDING (Free sample)"'], fill=fill_blue))
ws.conditional_formatting.add(work_range, CellIsRule(operator='equal', formula=['"FINISHED (Contact ready)"'], fill=fill_green))

for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

ws.column_dimensions['L'].width = 40

wb.save(output_file)
print(f"Process completed successfully!")
print(f"File generated: '{output_file}'.")